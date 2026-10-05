"""Read the baseline runs from Baselines.xlsx and turn every metric into a number.

One record per run (40 runs = 20 scenarios x 2 baselines). Cleaning rules:
  * Time to decision is typed as minutes:seconds but stored by the spreadsheet as hours:minutes,
    so 04:48 means 4 min 48 s.
  * A value recorded as a range ("20-25") becomes its midpoint.
  * Price cells of the form "$33-53 (whole bag) / $3-4.25 per-person amount" use the whole-bag
    amount, i.e. what the user has to pay at the store, which is what the other users recorded.
  * Runs are paired by position: each user's five Google rows and five DeepSeek rows are in the
    same scenario order (checked against the scenario text).
"""
import datetime
import re
from pathlib import Path

import openpyxl

XLSX = Path(__file__).resolve().parents[2] / "Baselines.xlsx"
FIRST_ROW, LAST_ROW = 3, 42

# key, column in the sheet, label, unit, +1 if higher is better / -1 if lower is better
METRICS = [
    ("shop_min", 7, "Grocery shopping time", "min", -1),
    ("prep_min", 8, "Meal preparation time", "min", -1),
    ("price_usd", 9, "Price", "US$", -1),
    ("tastiness", 10, "Tastiness", "0-5", +1),
    ("nutrition", 11, "Nutritional value", "Ofcom score", -1),
    ("diet_violation", 12, "Diet violation", "0/1", -1),
    ("decide_min", 6, "Time to decision", "min", -1),
]
USERS = {"koala": "User 1", "lion": "User 2", "kangaroo": "User 3", "hamster": "User 4"}
SHORT = {      # short labels for tables; the full scenario texts are in the sheet and in the writeup's appendix
    "koala": ["Asian", "Italian", "Green onion and ginger", "Fish", "Tasty and filling"],
    "lion": ["American, no beef", "Body building", "Bengali, healthy", "North Indian, high protein",
             "South Indian, high protein"],
    "kangaroo": ["Northern Chinese", "Southeast Chinese", "Chinese, more vegetables", "U.S.-style Chinese, tasty",
                 "Taiwanese"],
    "hamster": ["Mexican, high protein", "Korean vegetarian, kimchi", "Tofu and soy sauce", "Broccoli-based vegetarian",
                "Gnocchi"],
}
NUM = r"\d+(?:\.\d+)?"


def midpoint(text):
    """'20-25' -> 22.5, '30' -> 30.0"""
    nums = [float(x) for x in re.findall(NUM, text)]
    assert 1 <= len(nums) <= 2, text
    return sum(nums) / len(nums)


def to_number(key, value):
    if key == "decide_min":
        assert isinstance(value, datetime.time), value
        return value.hour + value.minute / 60          # stored h:m, meant min:s
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value)
    if key == "price_usd" and "whole bag" in text:
        text = text.split("(whole bag)")[0]
    return midpoint(text.replace("$", ""))


def load():
    ws = openpyxl.load_workbook(XLSX)["Main"]
    runs = []
    for r in range(FIRST_ROW, LAST_ROW + 1):
        method = "google" if str(ws.cell(r, 4).value).startswith("Google") else "deepseek"
        rec = dict(row=r, user=ws.cell(r, 2).value, method=method,
                   scenario_text=(ws.cell(r, 3).value or "").strip().strip('"“”'))
        rec["recorded"] = {}
        for key, col, *_ in METRICS:
            value = ws.cell(r, col).value
            rec[key] = to_number(key, value)
            rec["recorded"][key] = f"{value.hour}:{value.minute:02d}" if isinstance(value, datetime.time) else str(value)
        rec["recipe_text"] = re.sub(r"\s+", " ", ws.cell(r, 5).value or "").strip()
        runs.append(rec)

    pairs = []
    for user in USERS:
        g = [x for x in runs if x["user"] == user and x["method"] == "google"]
        d = [x for x in runs if x["user"] == user and x["method"] == "deepseek"]
        assert len(g) == len(d) == 5, user
        for i, (a, b) in enumerate(zip(g, d)):
            assert not b["scenario_text"] or a["scenario_text"] == b["scenario_text"], (a["row"], b["row"])
            b["scenario_text"] = a["scenario_text"]      # one cell is empty in the sheet; same scenario by position
            pairs.append(dict(user=user, user_id=USERS[user], index=i + 1, label=SHORT[user][i], google=a, deepseek=b))
    assert len(pairs) == 20
    return runs, pairs


if __name__ == "__main__":
    runs, pairs = load()
    for p in pairs:
        print(p["user"], p["index"], p["label"])
        for m in ("google", "deepseek"):
            print("   ", m.ljust(8), {k: round(p[m][k], 2) for k, *_ in METRICS})
