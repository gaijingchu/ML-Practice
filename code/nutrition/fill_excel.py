"""Write the Ofcom NPM scores into the 'Nutritional value' column (K) of the baselines spreadsheet.

Usage:
    python fill_excel.py [--xlsx ../../Baselines.xlsx]
Reads output/nutrition_scores.csv (run compute_nutrition.py first). Only column K is touched.
"""
import argparse
import csv
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
COLUMN = 11          # K: "Nutritional value"
HEADER_NOTE = "UK Ofcom Nutrient Profiling Model score per 100 g (lower = healthier; 4 or more = 'less healthy')"


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--xlsx", default=str(HERE.parent.parent / "Baselines.xlsx"))
    args = ap.parse_args()

    with open(HERE / "output" / "nutrition_scores.csv", newline="", encoding="utf-8") as fh:
        scores = {int(r["sheet_row"]): int(r["npm_score"]) for r in csv.DictReader(fh)}

    wb = openpyxl.load_workbook(args.xlsx)
    ws = wb["Main"]
    assert ws.cell(1, COLUMN).value == "Nutritional value", "column K is not the nutritional-value column"
    ws.cell(2, COLUMN).value = HEADER_NOTE
    for row, score in scores.items():
        assert ws.cell(row, 5).value, f"row {row} has no recipe text"
        ws.cell(row, COLUMN).value = score
    wb.save(args.xlsx)
    print(f"wrote {len(scores)} scores to column K of {args.xlsx}")


if __name__ == "__main__":
    main()
